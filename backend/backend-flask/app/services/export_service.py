"""
Servicio para exportación de datos
"""
import csv
import io
import json
from datetime import datetime, timedelta
from typing import List, Dict, Any, Optional, Union
from flask import current_app, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_LEFT

from app.repositories.sensor_repository import SensorRepository
from app.repositories.user_repository import UserRepository
from app.repositories.alert_repository import AlertRepository

class ExportService:
    """Servicio para exportación de datos en diferentes formatos"""
    
    def __init__(self, sensor_repository: SensorRepository, 
                 user_repository: UserRepository,
                 alert_repository: Optional[AlertRepository] = None):
        """
        Inicializar servicio
        
        Args:
            sensor_repository: Repository de sensores
            user_repository: Repository de usuarios
            alert_repository: Repository de alertas (opcional)
        """
        self.sensor_repository = sensor_repository
        self.user_repository = user_repository
        self.alert_repository = alert_repository
        current_app.logger.info("ExportService inicializado")
    
    def export_sensor_data(self, start_date: datetime, end_date: datetime, 
                          format_type: str = 'csv', filters: Optional[Dict] = None) -> Dict[str, Any]:
        """
        Exportar datos de sensores
        
        Args:
            start_date: Fecha de inicio
            end_date: Fecha de fin
            format_type: Tipo de formato (csv, excel, pdf, json)
            filters: Filtros adicionales
            
        Returns:
            Dict con el archivo generado
        """
        try:
            # Obtener datos de sensores
            sensor_data = self.sensor_repository.find_by_date_range(start_date, end_date, filters)
            
            if not sensor_data:
                return {
                    'success': False,
                    'message': 'No hay datos para exportar en el rango especificado'
                }
            
            # Generar archivo según el formato
            if format_type.lower() == 'csv':
                return self._generate_csv(sensor_data, 'sensor_data')
            elif format_type.lower() == 'excel':
                return self._generate_excel(sensor_data, 'sensor_data')
            elif format_type.lower() == 'pdf':
                return self._generate_pdf(sensor_data, 'sensor_data')
            elif format_type.lower() == 'json':
                return self._generate_json(sensor_data, 'sensor_data')
            else:
                return {
                    'success': False,
                    'message': f'Formato no soportado: {format_type}'
                }
                
        except Exception as e:
            current_app.logger.error(f"Error exportando datos de sensores: {e}")
            return {
                'success': False,
                'message': 'Error interno exportando datos'
            }
    
    def export_user_data(self, format_type: str = 'csv', 
                        filters: Optional[Dict] = None) -> Dict[str, Any]:
        """
        Exportar datos de usuarios
        
        Args:
            format_type: Tipo de formato
            filters: Filtros adicionales
            
        Returns:
            Dict con el archivo generado
        """
        try:
            # Obtener datos de usuarios
            user_data = self.user_repository.find_all(filters)
            
            if not user_data:
                return {
                    'success': False,
                    'message': 'No hay datos de usuarios para exportar'
                }
            
            # Limpiar datos sensibles
            cleaned_data = []
            for user in user_data:
                cleaned_user = {
                    'id': user.get('_id'),
                    'username': user.get('username'),
                    'email': user.get('email'),
                    'full_name': user.get('full_name'),
                    'role': user.get('role'),
                    'status': user.get('status'),
                    'location': user.get('location'),
                    'created_at': user.get('created_at'),
                    'last_login': user.get('last_login')
                }
                cleaned_data.append(cleaned_user)
            
            # Generar archivo según el formato
            if format_type.lower() == 'csv':
                return self._generate_csv(cleaned_data, 'user_data')
            elif format_type.lower() == 'excel':
                return self._generate_excel(cleaned_data, 'user_data')
            elif format_type.lower() == 'pdf':
                return self._generate_pdf(cleaned_data, 'user_data')
            elif format_type.lower() == 'json':
                return self._generate_json(cleaned_data, 'user_data')
            else:
                return {
                    'success': False,
                    'message': f'Formato no soportado: {format_type}'
                }
                
        except Exception as e:
            current_app.logger.error(f"Error exportando datos de usuarios: {e}")
            return {
                'success': False,
                'message': 'Error interno exportando datos'
            }
    
    def export_alert_data(self, start_date: datetime, end_date: datetime,
                         format_type: str = 'csv', filters: Optional[Dict] = None) -> Dict[str, Any]:
        """
        Exportar datos de alertas
        
        Args:
            start_date: Fecha de inicio
            end_date: Fecha de fin
            format_type: Tipo de formato
            filters: Filtros adicionales
            
        Returns:
            Dict con el archivo generado
        """
        try:
            if not self.alert_repository:
                return {
                    'success': False,
                    'message': 'Repository de alertas no disponible'
                }
            
            # Obtener datos de alertas
            alert_data = self.alert_repository.find_by_date_range(start_date, end_date, filters)
            
            if not alert_data:
                return {
                    'success': False,
                    'message': 'No hay alertas para exportar en el rango especificado'
                }
            
            # Generar archivo según el formato
            if format_type.lower() == 'csv':
                return self._generate_csv(alert_data, 'alert_data')
            elif format_type.lower() == 'excel':
                return self._generate_excel(alert_data, 'alert_data')
            elif format_type.lower() == 'pdf':
                return self._generate_pdf(alert_data, 'alert_data')
            elif format_type.lower() == 'json':
                return self._generate_json(alert_data, 'json_data')
            else:
                return {
                    'success': False,
                    'message': f'Formato no soportado: {format_type}'
                }
                
        except Exception as e:
            current_app.logger.error(f"Error exportando datos de alertas: {e}")
            return {
                'success': False,
                'message': 'Error interno exportando datos'
            }
    
    def export_system_report(self, start_date: datetime, end_date: datetime,
                           format_type: str = 'pdf') -> Dict[str, Any]:
        """
        Exportar reporte completo del sistema
        
        Args:
            start_date: Fecha de inicio
            end_date: Fecha de fin
            format_type: Tipo de formato
            
        Returns:
            Dict con el archivo generado
        """
        try:
            # Obtener estadísticas del sistema
            system_stats = self._generate_system_stats(start_date, end_date)
            
            # Generar reporte según el formato
            if format_type.lower() == 'pdf':
                return self._generate_system_report_pdf(system_stats, start_date, end_date)
            elif format_type.lower() == 'excel':
                return self._generate_system_report_excel(system_stats, start_date, end_date)
            else:
                return {
                    'success': False,
                    'message': f'Formato no soportado para reporte: {format_type}'
                }
                
        except Exception as e:
            current_app.logger.error(f"Error generando reporte del sistema: {e}")
            return {
                'success': False,
                'message': 'Error interno generando reporte'
            }
    
    def _generate_csv(self, data: List[Dict], filename_prefix: str) -> Dict[str, Any]:
        """Generar archivo CSV"""
        try:
            if not data:
                return {
                    'success': False,
                    'message': 'No hay datos para exportar'
                }
            
            # Crear buffer en memoria
            output = io.StringIO()
            
            # Obtener headers del primer elemento
            fieldnames = list(data[0].keys())
            
            # Escribir CSV
            writer = csv.DictWriter(output, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data)
            
            # Preparar respuesta
            output.seek(0)
            csv_content = output.getvalue()
            output.close()
            
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{filename_prefix}_{timestamp}.csv"
            
            return {
                'success': True,
                'filename': filename,
                'content': csv_content,
                'content_type': 'text/csv',
                'size': len(csv_content)
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando CSV: {e}")
            return {
                'success': False,
                'message': 'Error generando archivo CSV'
            }
    
    def _generate_excel(self, data: List[Dict], filename_prefix: str) -> Dict[str, Any]:
        """Generar archivo Excel"""
        try:
            if not data:
                return {
                    'success': False,
                    'message': 'No hay datos para exportar'
                }
            
            # Crear workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Datos"
            
            # Obtener headers
            headers = list(data[0].keys())
            
            # Configurar estilo del header
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            
            # Escribir headers
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # Escribir datos
            for row, record in enumerate(data, 2):
                for col, header in enumerate(headers, 1):
                    value = record.get(header, '')
                    # Formatear fechas
                    if isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    ws.cell(row=row, column=col, value=value)
            
            # Ajustar ancho de columnas
            for col, header in enumerate(headers, 1):
                column_letter = get_column_letter(col)
                ws.column_dimensions[column_letter].width = max(len(str(header)), 15)
            
            # Guardar en buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{filename_prefix}_{timestamp}.xlsx"
            
            return {
                'success': True,
                'filename': filename,
                'content': output.getvalue(),
                'content_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                'size': len(output.getvalue())
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando Excel: {e}")
            return {
                'success': False,
                'message': 'Error generando archivo Excel'
            }
    
    def _generate_pdf(self, data: List[Dict], filename_prefix: str) -> Dict[str, Any]:
        """Generar archivo PDF"""
        try:
            if not data:
                return {
                    'success': False,
                    'message': 'No hay datos para exportar'
                }
            
            # Crear buffer en memoria
            buffer = io.BytesIO()
            
            # Crear documento PDF
            doc = SimpleDocTemplate(buffer, pagesize=A4, 
                                  rightMargin=72, leftMargin=72,
                                  topMargin=72, bottomMargin=18)
            
            # Obtener estilos
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=16,
                spaceAfter=30,
                alignment=TA_CENTER
            )
            
            # Elementos del documento
            story = []
            
            # Título
            title = Paragraph("Reporte de Datos", title_style)
            story.append(title)
            story.append(Spacer(1, 12))
            
            # Información del reporte
            info_text = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            info_para = Paragraph(info_text, styles['Normal'])
            story.append(info_para)
            story.append(Spacer(1, 12))
            
            # Obtener headers
            headers = list(data[0].keys())
            
            # Preparar datos para la tabla
            table_data = [headers]
            for record in data[:100]:  # Limitar a 100 registros para PDF
                row = []
                for header in headers:
                    value = record.get(header, '')
                    if isinstance(value, datetime):
                        value = value.strftime('%Y-%m-%d %H:%M:%S')
                    row.append(str(value))
                table_data.append(row)
            
            # Crear tabla
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
            ]))
            
            story.append(table)
            
            # Construir PDF
            doc.build(story)
            
            # Preparar respuesta
            buffer.seek(0)
            pdf_content = buffer.getvalue()
            buffer.close()
            
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{filename_prefix}_{timestamp}.pdf"
            
            return {
                'success': True,
                'filename': filename,
                'content': pdf_content,
                'content_type': 'application/pdf',
                'size': len(pdf_content)
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando PDF: {e}")
            return {
                'success': False,
                'message': 'Error generando archivo PDF'
            }
    
    def _generate_json(self, data: List[Dict], filename_prefix: str) -> Dict[str, Any]:
        """Generar archivo JSON"""
        try:
            if not data:
                return {
                    'success': False,
                    'message': 'No hay datos para exportar'
                }
            
            # Preparar datos para JSON
            json_data = {
                'metadata': {
                    'generated_at': datetime.now().isoformat(),
                    'total_records': len(data),
                    'data_type': filename_prefix
                },
                'data': data
            }
            
            # Convertir a JSON
            json_content = json.dumps(json_data, indent=2, default=str, ensure_ascii=False)
            
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{filename_prefix}_{timestamp}.json"
            
            return {
                'success': True,
                'filename': filename,
                'content': json_content,
                'content_type': 'application/json',
                'size': len(json_content)
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando JSON: {e}")
            return {
                'success': False,
                'message': 'Error generando archivo JSON'
            }
    
    def _generate_system_stats(self, start_date: datetime, end_date: datetime) -> Dict[str, Any]:
        """Generar estadísticas del sistema"""
        try:
            # Estadísticas de sensores
            sensor_stats = self.sensor_repository.get_statistics_by_date_range(start_date, end_date)
            
            # Estadísticas de usuarios
            user_stats = self.user_repository.get_users_summary()
            
            # Estadísticas de alertas (si está disponible)
            alert_stats = {}
            if self.alert_repository:
                alert_stats = self.alert_repository.get_statistics_by_date_range(start_date, end_date)
            
            return {
                'period': {
                    'start_date': start_date.isoformat(),
                    'end_date': end_date.isoformat()
                },
                'sensor_statistics': sensor_stats,
                'user_statistics': user_stats,
                'alert_statistics': alert_stats,
                'generated_at': datetime.now().isoformat()
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando estadísticas: {e}")
            return {}
    
    def _generate_system_report_pdf(self, stats: Dict[str, Any], 
                                   start_date: datetime, end_date: datetime) -> Dict[str, Any]:
        """Generar reporte del sistema en PDF"""
        try:
            # Crear buffer en memoria
            buffer = io.BytesIO()
            
            # Crear documento PDF
            doc = SimpleDocTemplate(buffer, pagesize=A4,
                                  rightMargin=72, leftMargin=72,
                                  topMargin=72, bottomMargin=18)
            
            # Obtener estilos
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                spaceAfter=30,
                alignment=TA_CENTER
            )
            
            # Elementos del documento
            story = []
            
            # Título
            title = Paragraph("Reporte del Sistema de Monitoreo", title_style)
            story.append(title)
            story.append(Spacer(1, 12))
            
            # Información del período
            period_text = f"Período: {start_date.strftime('%Y-%m-%d')} a {end_date.strftime('%Y-%m-%d')}"
            period_para = Paragraph(period_text, styles['Heading2'])
            story.append(period_para)
            story.append(Spacer(1, 12))
            
            # Fecha de generación
            generated_text = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            generated_para = Paragraph(generated_text, styles['Normal'])
            story.append(generated_para)
            story.append(Spacer(1, 20))
            
            # Agregar estadísticas aquí...
            # (Implementación simplificada)
            
            # Construir PDF
            doc.build(story)
            
            # Preparar respuesta
            buffer.seek(0)
            pdf_content = buffer.getvalue()
            buffer.close()
            
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"system_report_{timestamp}.pdf"
            
            return {
                'success': True,
                'filename': filename,
                'content': pdf_content,
                'content_type': 'application/pdf',
                'size': len(pdf_content)
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando reporte PDF: {e}")
            return {
                'success': False,
                'message': 'Error generando reporte PDF'
            }
    
    def _generate_system_report_excel(self, stats: Dict[str, Any],
                                     start_date: datetime, end_date: datetime) -> Dict[str, Any]:
        """Generar reporte del sistema en Excel"""
        try:
            # Crear workbook
            wb = Workbook()
            
            # Hoja de resumen
            ws_summary = wb.active
            ws_summary.title = "Resumen"
            
            # Configurar estilo
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            
            # Título
            ws_summary['A1'] = "Reporte del Sistema de Monitoreo"
            ws_summary['A1'].font = Font(bold=True, size=16)
            ws_summary.merge_cells('A1:D1')
            
            # Información del período
            ws_summary['A3'] = f"Período: {start_date.strftime('%Y-%m-%d')} a {end_date.strftime('%Y-%m-%d')}"
            ws_summary['A4'] = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            
            # Agregar estadísticas aquí...
            # (Implementación simplificada)
            
            # Guardar en buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"system_report_{timestamp}.xlsx"
            
            return {
                'success': True,
                'filename': filename,
                'content': output.getvalue(),
                'content_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                'size': len(output.getvalue())
            }
            
        except Exception as e:
            current_app.logger.error(f"Error generando reporte Excel: {e}")
            return {
                'success': False,
                'message': 'Error generando reporte Excel'
            }
